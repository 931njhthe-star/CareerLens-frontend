"""No live API, database, credentials or external network needed."""

from email.message import Message
from pathlib import Path
from io import BytesIO
import json
import os
import tempfile
import unittest
from unittest.mock import patch
from urllib.error import URLError

import server


class UpstreamResponse:
    def __init__(self, status=200, headers=(), body=b'{"user":null}'):
        self.code = status
        self.headers = Message()
        self.headers["Content-Type"] = "application/json"
        for key, value in headers:
            self.headers[key] = value
        self.body = body

    def read(self):
        return self.body

    def __enter__(self):
        return self

    def __exit__(self, *_):
        return False


class Opener:
    def __init__(self, response=None, error=None):
        self.response = response or UpstreamResponse()
        self.error = error
        self.calls = []

    def open(self, request, timeout):
        self.calls.append((request, timeout))
        if self.error:
            raise self.error
        return self.response


class ServerTests(unittest.TestCase):
    def app(self, origin="", opener=None):
        with (
            patch.dict(os.environ, {}, clear=True),
            patch("server.build_opener", return_value=opener or Opener()),
        ):
            return server.create_app(origin)

    def test_origin_validation(self):
        for value in (
            "http://localhost:5101",
            "http://127.0.0.1:5101/",
            "http://[::1]:5101",
            "https://api.example.com",
            "https://127.0.0.1:8443",
        ):
            with self.subTest(value=value):
                self.assertTrue(server.validate_origin(value))
        for value in (
            "http://api.example.com",
            "ftp://localhost",
            "https://user:password@api.example.com",
            "https://api.example.com/api/v1",
            "https://api.example.com?key=value",
            "https://api.example.com/#secret",
            "https://api.example.com:0",
            "https://api.example.com:99999",
            "https://bad\\host",
            "https://bad%20host",
            "https://bad_host",
            "https://api.example.com\n",
            "//api.example.com",
            "https://",
        ):
            with self.subTest(value=value), self.assertRaises(ValueError):
                server.validate_origin(value)

    def test_unconfigured_and_unavailable_apis_are_errors(self):
        app = self.app()
        with app.test_client() as client:
            with client.get("/") as index:
                self.assertEqual(index.status_code, 200)
            self.assertFalse(client.get("/health").json["backend_configured"])
            response = client.get("/api/v1/auth/session")
            self.assertEqual(response.status_code, 503)
            self.assertEqual(response.json["error"]["code"], "backend_unconfigured")
        app = self.app("https://api.example.com", Opener(error=URLError("offline")))
        response = app.test_client().get("/api/v1/auth/session")
        self.assertEqual(response.status_code, 502)
        self.assertEqual(response.json["error"]["code"], "backend_unavailable")

    def test_proxy_keeps_csrf_origin_cookie_and_uses_api_host(self):
        opener = Opener()
        with self.app("https://api.example.com", opener).test_client() as client:
            client.set_cookie("careerlens_session", "test-session")
            response = client.post(
                "/api/v1/auth/login?case=1",
                json={"email": "student@example.com"},
                headers={"Origin": "http://localhost", "X-CSRF-Token": "test-csrf"},
            )
        self.assertEqual(response.status_code, 200)
        request, timeout = opener.calls[0]
        headers = {key.lower(): value for key, value in request.header_items()}
        self.assertEqual(request.full_url, "https://api.example.com/api/v1/auth/login?case=1")
        self.assertEqual(request.host, "api.example.com")
        self.assertNotIn("host", headers)
        self.assertEqual(headers["origin"], "http://localhost")
        self.assertEqual(headers["x-csrf-token"], "test-csrf")
        self.assertIn("careerlens_session=test-session", headers["cookie"])
        self.assertEqual(json.loads(request.data), {"email": "student@example.com"})
        self.assertEqual(timeout, 60)

    def test_proxy_preserves_backend_auth_idempotency_and_event_cursor(self):
        opener = Opener()
        response = self.app("http://127.0.0.1:5301", opener).test_client().post(
            "/api/v1/evaluations",
            json={"resume_id": "synthetic-resume", "job_posting_id": "synthetic-job"},
            headers={"Authorization": "Bearer synthetic-token", "Idempotency-Key": "synthetic-run", "Last-Event-ID": "12", "X-Untrusted": "drop"},
        )
        self.assertEqual(response.status_code, 200)
        headers = {key.lower(): value for key, value in opener.calls[0][0].header_items()}
        self.assertEqual(headers["authorization"], "Bearer synthetic-token")
        self.assertEqual(headers["idempotency-key"], "synthetic-run")
        self.assertEqual(headers["last-event-id"], "12")
        self.assertNotIn("x-untrusted", headers)

    def test_local_catalog_reads_only_numbered_top_level_markdown(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            for name in ("010.md", "002.md", "README.md", "00_목록.md", "001.txt"):
                (root / name).write_text("가상 공고 " + name, encoding="utf-8-sig")
            (root / "nested").mkdir()
            (root / "nested" / "003.md").write_text("excluded", encoding="utf-8")
            opener = Opener(error=AssertionError("Local catalog must not contact API"))
            app = self.app("http://127.0.0.1:5301", opener)
            with patch.dict(os.environ, {"CAREERLENS_JOB_POSTINGS_DIR": directory}):
                response = app.test_client().get("/api/v1/local-data/career-markdown?path=ignored")
            self.assertEqual(response.status_code, 200)
            self.assertEqual([item["id"] for item in response.json["items"]], ["002", "010"])
            self.assertEqual(response.json["items"][0]["content"], "가상 공고 002.md")
            self.assertEqual(opener.calls, [])

    def test_local_catalog_fails_without_exposing_local_paths(self):
        with tempfile.TemporaryDirectory() as directory:
            with patch("server.career_postings_directory", return_value=Path(directory)):
                response = self.app().test_client().get("/api/v1/local-data/career-markdown")
            self.assertEqual(response.status_code, 503)
            self.assertEqual(response.json["error"]["code"], "career_markdown_unavailable")
            self.assertNotIn(directory, response.get_data(as_text=True))

    def test_resume_conversion_handles_markdown_and_legacy_korean_text_without_upstream(self):
        opener = Opener(error=AssertionError("Conversion must not contact API"))
        app = self.app("http://127.0.0.1:5301", opener)
        for filename, content in (
            ("resume.md", "# 가상 이력서\r\n테스트 경력".encode("utf-8-sig")),
            ("C:\\uploads\\이력서.TXT", "가상 이력서\r\n테스트 경력".encode("cp949")),
        ):
            with self.subTest(filename=filename):
                response = app.test_client().post("/api/v1/local-data/convert-resume", data={"file": (BytesIO(content), filename)})
                self.assertEqual(response.status_code, 200)
                self.assertTrue(response.json["filename"].endswith(".md"))
                self.assertNotIn("\\", response.json["filename"])
                self.assertIn("가상 이력서\n테스트 경력", response.json["text"])
                self.assertEqual(response.headers["Cache-Control"], "no-store")
        self.assertEqual(opener.calls, [])

    def test_resume_conversion_extracts_docx_paragraphs_and_tables_in_order(self):
        from docx import Document
        document = Document()
        document.add_paragraph("가상 이력서")
        table = document.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "분석"
        table.cell(0, 1).text = "테스트 경력"
        document.add_paragraph("마지막")
        content = BytesIO()
        document.save(content)
        content.seek(0)
        response = self.app().test_client().post("/api/v1/local-data/convert-resume", data={"file": (content, "synthetic.docx")})
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json, {"filename": "synthetic.md", "text": "가상 이력서\n분석\t테스트 경력\n마지막"})

    def test_resume_conversion_extracts_pdf_text(self):
        from pypdf import PdfWriter
        from pypdf.generic import DictionaryObject, NameObject, DecodedStreamObject
        writer = PdfWriter()
        page = writer.add_blank_page(width=400, height=200)
        font = DictionaryObject({NameObject("/Type"): NameObject("/Font"), NameObject("/Subtype"): NameObject("/Type1"), NameObject("/BaseFont"): NameObject("/Helvetica")})
        page[NameObject("/Resources")] = DictionaryObject({NameObject("/Font"): DictionaryObject({NameObject("/F1"): writer._add_object(font)})})
        content = DecodedStreamObject()
        content.set_data(b"BT /F1 12 Tf 20 100 Td (Synthetic resume experience) Tj ET")
        page[NameObject("/Contents")] = writer._add_object(content)
        uploaded = BytesIO()
        writer.write(uploaded)
        uploaded.seek(0)
        response = self.app().test_client().post("/api/v1/local-data/convert-resume", data={"file": (uploaded, "synthetic.pdf")})
        self.assertEqual(response.status_code, 200)
        self.assertIn("Synthetic resume experience", response.json["text"])

    def test_resume_conversion_rejects_missing_unsupported_empty_and_disguised_binary(self):
        client = self.app().test_client()
        response = client.post("/api/v1/local-data/convert-resume")
        self.assertEqual(response.status_code, 400)
        self.assertEqual(response.json["error"]["code"], "file_required")
        for filename, data, status in (("test.exe", b"MZexample", 415), ("empty.md", b"", 400), ("fake.txt", b"%PDF-1.7", 400), ("fake.docx", b"not a docx", 400), ("bad.md", b"text\x00data", 400)):
            with self.subTest(filename=filename):
                response = client.post("/api/v1/local-data/convert-resume", data={"file": (BytesIO(data), filename)})
                self.assertEqual(response.status_code, status)
                self.assertIn("message", response.json["error"])

    def test_resume_conversion_enforces_file_limit_before_parsing(self):
        with patch("server.extract_resume", side_effect=AssertionError("Oversized data must not reach parser")):
            response = self.app().test_client().post("/api/v1/local-data/convert-resume", data={"file": (BytesIO(b"x" * (server.MAX_UPLOAD_BYTES + 1)), "large.txt")})
        self.assertEqual(response.status_code, 413)
        self.assertEqual(response.json["error"]["code"], "file_too_large")
        response.request.input_stream.close()
        response.close()

    def test_cookie_domain_removed_security_attributes_and_multiple_headers_preserved(self):
        upstream = UpstreamResponse(
            headers=[
                (
                    "Set-Cookie",
                    "session=one; Domain=api.example.com; Path=/; Secure; HttpOnly; SameSite=Lax",
                ),
                ("Set-Cookie", "state=two; DOMAIN=.example.com; Path=/api; Secure; SameSite=None"),
            ]
        )
        response = (
            self.app("https://api.example.com", Opener(upstream))
            .test_client()
            .get("/api/v1/auth/session")
        )
        cookies = response.headers.getlist("Set-Cookie")
        self.assertEqual(len(cookies), 2)
        self.assertEqual(cookies[0], "session=one; Path=/; Secure; HttpOnly; SameSite=Lax")
        self.assertEqual(cookies[1], "state=two; Path=/api; Secure; SameSite=None")

    def test_api_redirects_stay_local_and_oauth_provider_redirect_is_not_followed(self):
        for destination, expected in (
            (
                "https://api.example.com/api/v1/auth/callback?code=test",
                "/api/v1/auth/callback?code=test",
            ),
            ("https://api.example.com:443/?auth_error=test", "/?auth_error=test"),
            ("/?reset_token=test", "/?reset_token=test"),
            (
                "https://accounts.google.com/o/oauth2/auth?state=test",
                "https://accounts.google.com/o/oauth2/auth?state=test",
            ),
            ("https://api.example.com.evil.test/", "https://api.example.com.evil.test/"),
        ):
            with self.subTest(destination=destination):
                opener = Opener(UpstreamResponse(302, [("Location", destination)]))
                response = (
                    self.app("https://api.example.com", opener)
                    .test_client()
                    .get("/api/v1/auth/oauth/google")
                )
                self.assertEqual(response.status_code, 302)
                self.assertEqual(response.headers["Location"], expected)
                self.assertEqual(len(opener.calls), 1)
        self.assertIsNone(
            server.NoRedirect().redirect_request(
                None, None, 302, "", {}, "https://accounts.google.com"
            )
        )

    def test_env_file_is_not_executed_and_external_environment_wins(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / ".env"
            path.write_text(
                '# team settings\nBACKEND_URL="https://api.example.com"\nFRONTEND_PORT=5100\nLITERAL=$(not-a-command)\n',
                encoding="utf-8",
            )
            with patch.dict(os.environ, {"BACKEND_URL": "http://127.0.0.1:7777"}, clear=True):
                server.load_local_env(path)
                self.assertEqual(os.environ["BACKEND_URL"], "http://127.0.0.1:7777")
                self.assertEqual(os.environ["FRONTEND_PORT"], "5100")
                self.assertEqual(os.environ["LITERAL"], "$(not-a-command)")

    def test_explicit_argument_overrides_environment_and_static_files_do_not_expose_env(self):
        opener = Opener()
        with (
            patch.dict(os.environ, {"BACKEND_URL": "https://ignored.example.com"}, clear=True),
            patch("server.build_opener", return_value=opener),
        ):
            app = server.create_app("http://127.0.0.1:7777")
        client = app.test_client()
        client.get("/api/v1/auth/session")
        self.assertTrue(opener.calls[0][0].full_url.startswith("http://127.0.0.1:7777/"))
        self.assertEqual(client.get("/.env").status_code, 404)
        self.assertEqual(client.get("/src/.env").status_code, 404)
        with client.get("/favicon.svg") as favicon:
            self.assertEqual(favicon.status_code, 200)

    def test_matching_bundle_is_local_and_works_with_existing_csp(self):
        with self.app().test_client() as client:
            with client.get("/src/features/analysis/matching-eye.js") as response:
                self.assertEqual(response.status_code, 200)
                self.assertIn(response.mimetype, {"text/javascript", "application/javascript"})
                self.assertIn(b"mountMatchingEye", response.data)
                policy = response.headers["Content-Security-Policy"]
                self.assertIn("script-src 'self'", policy)
                self.assertIn("style-src 'self' 'unsafe-inline'", policy)
                self.assertNotIn("'unsafe-eval'", policy)
                self.assertNotIn("https://", policy)

    def test_status_radar_modules_are_served_as_local_javascript(self):
        with self.app().test_client() as client:
            for name in ("status-radar.js", "status-radar-data.js"):
                with self.subTest(module=name):
                    with client.get("/src/features/analysis/" + name) as response:
                        self.assertEqual(response.status_code, 200)
                        self.assertIn(
                            response.mimetype, {"text/javascript", "application/javascript"}
                        )
                        self.assertIn(
                            "script-src 'self'", response.headers["Content-Security-Policy"]
                        )

    def test_design_lab_and_its_local_assets_work_without_contacting_the_configured_api(self):
        opener = Opener(error=AssertionError("The design lab must not contact an upstream API"))
        with self.app("https://api.example.com", opener).test_client() as client:
            with client.get("/") as index:
                expected_policy = index.headers["Content-Security-Policy"]
            for path in (
                "/design-lab",
                "/design-lab?theme=cobalt&layout=studio&eye=glacier&pyramid=ice&effect=comet",
                "/design-lab?theme=%3Cscript%3Ealert(1)%3C%2Fscript%3E&backend=https://invalid.test",
            ):
                with self.subTest(path=path), client.get(path) as response:
                    self.assertEqual(response.status_code, 200)
                    self.assertEqual(response.mimetype, "text/html")
                    self.assertEqual(
                        response.data, (server.ROOT / "public" / "design-lab.html").read_bytes()
                    )
                    self.assertIn(b'/src/features/design-lab/index.js', response.data)
                    self.assertNotIn(b'/src/app/main.js', response.data)
                    self.assertNotIn(b"<script>alert(1)</script>", response.data)
                    self.assertEqual(response.headers["Content-Security-Policy"], expected_policy)
                    self.assertEqual(response.headers["X-Content-Type-Options"], "nosniff")
                    self.assertEqual(response.headers["X-Frame-Options"], "DENY")
                    self.assertEqual(response.headers["Cache-Control"], "no-store")
            for name in ("index.js", "presets.js", "preview.js", "transition.js", "lab.css", "preview.css"):
                with self.subTest(asset=name):
                    with client.get("/src/features/design-lab/" + name) as response:
                        self.assertEqual(response.status_code, 200)
                        expected_types = (
                            {"text/javascript", "application/javascript"}
                            if name.endswith(".js")
                            else {"text/css"}
                        )
                        self.assertIn(response.mimetype, expected_types)
                        self.assertEqual(
                            response.headers["Content-Security-Policy"], expected_policy
                        )
                        self.assertNotIn("'unsafe-eval'", expected_policy)
                        self.assertNotIn("https://", expected_policy)
        self.assertEqual(opener.calls, [])

    def test_design_lab_does_not_make_environment_or_parent_paths_public(self):
        with self.app().test_client() as client:
            for path in (
                "/.env",
                "/design-lab/.env",
                "/design-lab/../.env",
                "/src/../.env",
                "/src/features/design-lab/../../../.env",
                "/src/features/design-lab/%2E%2E/%2E%2E/%2E%2E/.env",
                "/public/../.env",
                "/public/design-lab.html",
            ):
                with self.subTest(path=path), client.get(path) as response:
                    self.assertEqual(response.status_code, 404)
                    self.assertIn("script-src 'self'", response.headers["Content-Security-Policy"])


if __name__ == "__main__":
    unittest.main()
